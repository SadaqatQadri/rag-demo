"""
loaders.py — turn whatever the client sends you into plain text.

Clients do not send tidy markdown. They send a 40-page PDF export, three Word
docs, and a link to their help centre. This module is the front door: it reads
each supported format, returns plain text, and tells you loudly when a file
gave back nothing usable (almost always a scanned PDF that needs OCR).

It also counts words, because "up to 20 pages" is meaningless across formats —
scope is quoted in words and this is what measures it.
"""

import os
import re
import glob

SUPPORTED = (".md", ".txt", ".pdf", ".docx")

# One "page" of scope = this many words. Keep in sync with what the gig says.
WORDS_PER_PAGE = 500

# Below this many words, a PDF page is almost certainly a scan with no text layer.
SCANNED_PDF_WORDS_PER_PAGE = 10


class LoadError(Exception):
    pass


def _clean(text: str) -> str:
    """Normalise whitespace without destroying paragraph breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\xa0", " ")           # non-breaking spaces from Word/web
    text = re.sub(r"[ \t]+", " ", text)        # collapse runs of spaces
    text = re.sub(r" *\n *", "\n", text)       # trim around newlines
    text = re.sub(r"\n{3,}", "\n\n", text)     # cap blank runs at one blank line
    return text.strip()


def load_text_file(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return _clean(f.read())


def load_pdf(path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise LoadError("pypdf is not installed — run: pip install pypdf")

    try:
        reader = PdfReader(path)
    except Exception as e:
        raise LoadError(f"could not open PDF ({e})")

    if reader.is_encrypted:
        try:
            reader.decrypt("")          # many PDFs are 'encrypted' with an empty password
        except Exception:
            raise LoadError("PDF is password protected — ask the client for an unlocked copy")

    pages = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")            # one bad page shouldn't kill the whole file

    text = _clean("\n\n".join(pages))
    words = len(text.split())
    if reader.pages and words / max(len(reader.pages), 1) < SCANNED_PDF_WORDS_PER_PAGE:
        raise LoadError(
            f"almost no extractable text ({words} words across {len(reader.pages)} pages) — "
            "this is probably a scanned PDF and needs OCR before it can be indexed"
        )
    return text


def load_docx(path: str) -> str:
    try:
        import docx
    except ImportError:
        raise LoadError("python-docx is not installed — run: pip install python-docx")

    try:
        d = docx.Document(path)
    except Exception as e:
        raise LoadError(f"could not open Word file ({e})")

    parts = [p.text for p in d.paragraphs]

    # Tables hold a lot of the real answers in policy documents, so don't skip them.
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return _clean("\n\n".join(p for p in parts if p.strip()))


LOADERS = {
    ".md": load_text_file,
    ".txt": load_text_file,
    ".pdf": load_pdf,
    ".docx": load_docx,
}


def load_file(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext not in LOADERS:
        raise LoadError(f"unsupported file type '{ext}' (supported: {', '.join(SUPPORTED)})")
    text = LOADERS[ext](path)
    if not text.strip():
        raise LoadError("file contained no readable text")
    return text


def word_count(text: str) -> int:
    return len(text.split())


def pages_equivalent(words: int) -> float:
    return round(words / WORDS_PER_PAGE, 1)


def load_directory(docs_dir: str):
    """Load every supported file in a folder.

    Returns (documents, problems) where documents is a list of
    {source, text, words} and problems is a list of {source, error}.
    Nothing raises — a client folder with one broken file should still index
    the rest, and you should see exactly which one failed and why.
    """
    documents, problems = [], []
    paths = []
    for ext in SUPPORTED:
        paths.extend(glob.glob(os.path.join(docs_dir, f"*{ext}")))

    for path in sorted(paths):
        source = os.path.basename(path)
        try:
            text = load_file(path)
        except LoadError as e:
            problems.append({"source": source, "error": str(e)})
            continue
        except Exception as e:  # anything unexpected, still keep going
            problems.append({"source": source, "error": f"unexpected error: {e}"})
            continue
        documents.append({"source": source, "text": text, "words": word_count(text)})

    return documents, problems


def scope_report(documents, problems) -> str:
    """A short human-readable summary. Run this before quoting a client."""
    lines = []
    total = 0
    for d in documents:
        total += d["words"]
        lines.append(f"  {d['source']:<40} {d['words']:>7,} words  (~{pages_equivalent(d['words'])} pages)")
    if problems:
        lines.append("")
        for p in problems:
            lines.append(f"  SKIPPED {p['source']}: {p['error']}")
    header = (
        f"{len(documents)} file(s) loaded, {total:,} words total "
        f"(~{pages_equivalent(total)} pages at {WORDS_PER_PAGE} words/page)"
    )
    return header + "\n" + "\n".join(lines)
